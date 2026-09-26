#!/usr/bin/env python3
"""Generate the demo repository's test suite and the QA inventory.

Why a generator: the fixture has to be *exactly* 500 collected tests, with 41 of
them structurally reachable from the changed module, 6 in the representation-
coupled consumer, and 75 in the inertly-changed billing module. Maintaining that
by hand is how a fixture drifts. The generator is checked in as provenance; its
output is checked in as the repository.

Independent-oracle rule: every expected value below is computed by a *reference
implementation in this file*, never by calling the code under test. A generator
that imported the app would make the fixture tautological.

The exact per-block counts are load-bearing: they are what makes the documented
demo numbers (500 / 47 / 41+6 / 11 modules / 18 direct / 23 transitive) come out.
Changing a case count changes the artefact, so the total is asserted at the end.

Run:  python3 tools/build_demo_suite.py
"""

import csv
import hashlib
import pathlib

ROOT = pathlib.Path(__file__).resolve().parents[1]
DEMO = ROOT / "demo_repo"
TESTS = DEMO / "tests"

EXPECTED_TOTAL = 500

# The report_worker block is hand-written (it contains the three contract tests
# and the subprocess-driven producer), so only its inventory rows live here.
# Order must match the parametrize ids in tools/templates/test_report_worker.py.
REPORT_WORKER_ROWS = [
    ("T-0338", "test_parse_keeps_unknown_payloads", "parsing: payload that is not three fields is kept", 22),
    ("T-0339", "test_parse_pipe_records", "parsing: a three-field record parses field by field", 19),
    ("T-0340", "test_recent_entries_round_trip_contract", "contract: every written field survives the round trip", 31),
    ("T-0341", "test_recent_entries_round_trip_statuses", "contract: the status field survives, in order", 28),
    ("T-0342", "test_recent_entries_round_trip_totals", "contract: the counts written are the counts reported", 7),
    ("T-0343", "test_summarise_entries_totals_counts", "summarise: counts add up across records", 15),
]


# --------------------------------------------------------------------------- #
# reference implementations (the fixture's independent oracle)                 #
# --------------------------------------------------------------------------- #
def ref_slug(text):
    out = []
    for char in text.strip().lower():
        if char.isalnum():
            out.append(char)
        elif char in " -_/.":
            out.append("-")
    collapsed = []
    for char in out:
        if char == "-" and collapsed and collapsed[-1] == "-":
            continue
        collapsed.append(char)
    return "".join(collapsed).strip("-")


def ref_truncate(text, limit):
    if limit < 0:
        raise ValueError("limit must not be negative")
    if len(text) <= limit:
        return text
    if limit <= 3:
        return "." * limit
    return text[: limit - 3] + "..."


_UNITS_MS = {"ms": 1, "s": 1000, "m": 60000, "h": 3600000}


def ref_parse_duration_ms(text):
    for suffix in ("ms", "h", "m", "s"):
        if text.endswith(suffix):
            quantity = text[: -len(suffix)]
            if not quantity.isdigit():
                raise ValueError(f"not a duration: {text!r}")
            return int(quantity) * _UNITS_MS[suffix]
    raise ValueError(f"not a duration: {text!r}")


def ref_format_duration(ms):
    for unit in ("h", "m", "s"):
        size = _UNITS_MS[unit]
        if ms and ms % size == 0:
            return f"{ms // size}{unit}"
    return f"{ms}ms"


def ref_is_valid_id(value):
    if not isinstance(value, str):
        return False
    for prefix in ("ACCT", "INV", "TXN"):
        if value.startswith(f"{prefix}-"):
            digits = value[len(prefix) + 1 :]
            return len(digits) == 6 and digits.isdigit()
    return False


def ref_next_sequence(values):
    highest = 0
    for value in values:
        if ref_is_valid_id(value):
            highest = max(highest, int(value.split("-")[-1]))
    return highest + 1


def ref_to_cents(amount_text):
    text = amount_text.strip()
    sign = 1
    if text[:1] in "+-":
        if text[0] == "-":
            sign = -1
        text = text[1:]
    if not text:
        raise ValueError(f"not an amount: {amount_text!r}")
    whole, dot, fraction = text.partition(".")
    if not whole.isdigit() or (dot and (not fraction.isdigit() or len(fraction) > 2)):
        raise ValueError(f"not an amount: {amount_text!r}")
    cents = int(whole) * 100
    if dot:
        cents += int(fraction.ljust(2, "0"))
    return sign * cents


def ref_format_amount(cents):
    sign = "-" if cents < 0 else ""
    magnitude = abs(cents)
    return f"{sign}{magnitude // 100}.{magnitude % 100:02d}"


def ref_apply_rate(cents, basis_points):
    return cents * basis_points // 10000


def ref_split_amount(cents, parts):
    if parts <= 0:
        raise ValueError("parts must be positive")
    base = cents // parts
    shares = [base] * parts
    shares[0] += cents - base * parts
    return shares


def ref_hash_token(token):
    return hashlib.sha256(token.encode("utf-8")).hexdigest()[:16]


ROLE_SCOPES = {"viewer": ("read",), "operator": ("read", "write"), "admin": ("read", "write", "delete")}


def ref_scopes_for(role):
    return sorted(ROLE_SCOPES.get(role, ()))


def ref_is_allowed(role, scope):
    return scope in ROLE_SCOPES.get(role, ())


def ref_load_config(overrides=None):
    merged = {"region": "eu-west", "retries": 3, "timeout_s": 30, "verbose": False}
    if overrides:
        for key, value in overrides.items():
            if key not in merged:
                raise KeyError(f"unknown configuration key: {key}")
            if not isinstance(value, type(merged[key])):
                raise TypeError(f"configuration key {key} expects {type(merged[key]).__name__}")
            merged[key] = value
    return merged


def ref_describe(config):
    return ",".join(f"{key}={config[key]}" for key in sorted(config))


def ref_apply_fee(amount_cents):
    return amount_cents * 25 // 10000


def ref_build_receipt(account_id, amount_cents):
    fee = ref_apply_fee(amount_cents)
    return {
        "account_id": account_id,
        "amount_cents": amount_cents,
        "fee_cents": fee,
        "net_cents": amount_cents - fee,
        "status": "settled",
    }


def ref_make_invoice(account_id, line_rows, tax_bps=0):
    subtotal = sum(row["unit_cents"] * row["quantity"] for row in line_rows)
    tax = subtotal * tax_bps // 10000
    return {
        "account_id": account_id,
        "lines": list(line_rows),
        "subtotal_cents": subtotal,
        "tax_cents": tax,
        "total_cents": subtotal + tax,
    }


def ref_base_price(sku):
    return {"starter": 900, "standard": 2500, "premium": 7900}[sku]


def ref_discounted(price_cents, percent):
    return price_cents - price_cents * percent // 100


def ref_tier_for(units):
    for threshold, name in ((100, "volume"), (10, "team"), (1, "single")):
        if units >= threshold:
            return name
    return "none"


def ref_discount_valid(code):
    if not isinstance(code, str) or len(code) != 6:
        return False
    return code[:4] == "SAVE" and all(char.isdigit() and char.isascii() for char in code[4:])


def ref_discount_for(code, cents):
    return cents * int(code[4:]) // 100 if ref_discount_valid(code) else 0


# --------------------------------------------------------------------------- #
# blocks: one per test file, counts pinned                                    #
# --------------------------------------------------------------------------- #
def generate_all():
    blocks = []

    def block(filename, module, doc, imports, functions, template=None, rows=None):
        blocks.append({
            "file": filename,
            "module": module,
            "doc": doc,
            "imports": imports,
            "functions": functions,
            "template": template,
            "rows": rows,
        })

    def fn(name, params, cases, body, fixtures=()):
        """A parametrized test: param names, case tuples, fixtures, and a body."""
        return {"name": name, "params": list(params), "cases": cases, "body": body, "fixtures": list(fixtures)}

    # ---- app/config.py : 30 ------------------------------------------------
    overrides = [
        {},
        {"region": "us-east"},
        {"retries": 0},
        {"timeout_s": 1},
        {"verbose": True},
        {"region": "ap-south", "retries": 5},
        {"retries": 9, "timeout_s": 90},
        {"verbose": False, "region": "sa-east"},
    ]
    configs = [
        {"region": "eu-west", "retries": 3, "timeout_s": 30, "verbose": False},
        {"region": "us-east", "retries": 0, "timeout_s": 1, "verbose": True},
        {"region": "ap-south", "retries": 9, "timeout_s": 90, "verbose": False},
        {"region": "sa-east", "retries": 1, "timeout_s": 10, "verbose": True},
        {"region": "r1", "retries": 1, "timeout_s": 10, "verbose": True},
        {"region": "r2", "retries": 2, "timeout_s": 20, "verbose": False},
        {"region": "r3", "retries": 3, "timeout_s": 30, "verbose": True},
        {"region": "r4", "retries": 4, "timeout_s": 40, "verbose": False},
    ]
    block(
        "test_config.py",
        "app.config",
        "Configuration merging and rendering.",
        "from app.config import DEFAULTS, describe, is_valid_key, load_config",
        [
            fn("test_merge_defaults", ["overrides", "expected"],
               [(o, ref_load_config(o)) for o in overrides],
               "assert load_config(overrides) == expected"),
            fn("test_is_valid_key", ["key", "expected"],
               [("region", True), ("retries", True), ("timeout_s", True), ("verbose", True),
                ("unknown", False), ("", False), (None, False), (7, False), ("Region", False), ("region ", False)],
               "assert is_valid_key(key) is expected"),
            fn("test_describe", ["config", "expected"],
               [(c, ref_describe(c)) for c in configs],
               "assert describe(config) == expected"),
            fn("test_load_config_rejects_bad_overrides", ["overrides", "expected_error"],
               [({"nosuchkey": 1}, "KeyError"), ({"RETRIES": 2}, "KeyError"),
                ({"region": 5}, "TypeError"), ({"verbose": "yes"}, "TypeError")],
               "errors = {'KeyError': KeyError, 'TypeError': TypeError}\n"
               "with pytest.raises(errors[expected_error]):\n"
               "    load_config(overrides)"),
        ],
    )

    # ---- app/models.py : 40 ------------------------------------------------
    lines = [
        {"unit_cents": 100, "quantity": 3}, {"unit_cents": 0, "quantity": 5},
        {"unit_cents": 250, "quantity": 1}, {"unit_cents": 999, "quantity": 7},
        {"unit_cents": 1, "quantity": 1}, {"unit_cents": 1250, "quantity": 2},
        {"unit_cents": 40, "quantity": 10}, {"unit_cents": 7, "quantity": 9},
        {"unit_cents": 5000, "quantity": 3}, {"unit_cents": 33, "quantity": 33},
    ]
    invoice_inputs = [
        ("acct", [{"unit_cents": 100, "quantity": 2}], 0),
        ("acct", [{"unit_cents": 100, "quantity": 2}], 1900),
        ("acct", [], 0),
        ("acct", [{"unit_cents": 2500, "quantity": 1}], 500),
        ("acct", [{"unit_cents": 999, "quantity": 3}], 2500),
        ("acct", [{"unit_cents": 1, "quantity": 1}], 1),
        ("acct", [{"unit_cents": 40, "quantity": 10}, {"unit_cents": 60, "quantity": 10}], 1000),
        ("acct", [{"unit_cents": 0, "quantity": 10}], 2000),
        ("acct", [{"unit_cents": 123456, "quantity": 2}], 725),
        ("acct", [{"unit_cents": 5, "quantity": 5}], 9999),
    ]
    block(
        "test_models.py",
        "app.models",
        "Domain record construction.",
        "from app.models import invoice_total, line_total, make_account, make_invoice",
        [
            fn("test_make_account", ["account_id", "owner", "expected"],
               [(f"ACCT-{i:06d}", f"owner-{i}",
                 {"id": f"ACCT-{i:06d}", "owner": f"owner-{i}", "balance_cents": 0, "state": "open"})
                for i in range(10)],
               "assert make_account(account_id, owner) == expected"),
            fn("test_line_total", ["line", "expected"],
               [(row, row["unit_cents"] * row["quantity"]) for row in lines],
               "assert line_total(line) == expected"),
            fn("test_make_invoice", ["account_id", "line_rows", "tax_bps", "expected"],
               [(a, lr, t, ref_make_invoice(a, lr, t)) for a, lr, t in invoice_inputs],
               "assert make_invoice(account_id, line_rows, tax_bps) == expected"),
            fn("test_invoice_total", ["invoice", "expected"],
               [(ref_make_invoice(a, lr, t),
                 ref_make_invoice(a, lr, t)["subtotal_cents"] + ref_make_invoice(a, lr, t)["tax_cents"])
                for a, lr, t in invoice_inputs],
               "assert invoice_total(invoice) == expected"),
        ],
    )

    # ---- app/services/billing_service.py : 75 ------------------------------
    amounts = [0, 1, 39, 40, 100, 101, 999, 1000, 2500, 10000, 12345, 99999, 100000, 250000,
               400, 800, 1200, 1600, 2000, 2400, 3600, 4800, 6000, 7200, 8400]
    block(
        "test_billing_service.py",
        "app.services.billing_service",
        "Fee arithmetic and receipts.",
        "from app.services.billing_service import apply_fee, build_receipt, settle",
        [
            fn("test_apply_fee", ["amount_cents", "expected"],
               [(amount, ref_apply_fee(amount)) for amount in amounts],
               "assert apply_fee(amount_cents) == expected"),
            fn("test_build_receipt", ["account_id", "amount_cents", "expected"],
               [(f"ACCT-{amount:06d}", amount, ref_build_receipt(f"ACCT-{amount:06d}", amount))
                for amount in amounts],
               "assert build_receipt(account_id, amount_cents) == expected"),
            fn("test_settle", ["account_id", "amount_cents", "expected"],
               [(f"ACCT-{amount:06d}", amount, ref_build_receipt(f"ACCT-{amount:06d}", amount))
                for amount in amounts],
               "assert settle(account_id, amount_cents) == expected"),
        ],
    )

    # ---- app/utils/text.py : 60 -------------------------------------------
    slug_inputs = [
        "  Mixed Case 1/Part_1  ", "Invoice Service", "ALREADY-SLUGGED", "dots.and.dots",
        "under_scores_here", "trailing---", "---leading", "a b  c   d", "Report/Worker", "",
        "   ", "9lives", "MixedCASE99", "slashes//double", "spaces and-dashes_and/slashes",
        "  spaced  ", "one", "Two Words", "three-words-here", "x/y/z",
    ]
    word_inputs = [
        "", " ", "one", "one two", "one  two   three", "a b c d e", "\ttab\tseparated\t",
        "trailing ", " leading", "many many many many many", "x", "x y", "x y z", "x y z w",
        "x y z w v", "line\nbreak", "line\n\nbreak", "a-b c-d", "  spaced  out  ", "unicode café count",
    ]
    truncate_inputs = [
        ("abcdef", 3), ("abcdef", 0), ("abcdef", 4), ("abcdef", 6), ("abcdef", 10),
        ("", 0), ("abc", 2), ("abc", 1), ("a much longer sentence to truncate", 12),
        ("a much longer sentence to truncate", 13), ("twenty-five characters", 20),
        ("twenty-five characters", 25), ("twenty-five characters", 26), ("short", 5),
        ("short", 100), ("0123456789", 8), ("0123456789", 3), ("0123456789", 2),
        ("unicode café", 6), ("unicode café", 4),
    ]
    block(
        "test_utils_text.py",
        "app.utils.text",
        "Slugging, word counting and truncation.",
        "from app.utils.text import normalise_slug, truncate, word_count",
        [
            fn("test_normalise_slug", ["text", "expected"],
               [(t, ref_slug(t)) for t in slug_inputs],
               "assert normalise_slug(text) == expected"),
            fn("test_word_count", ["text", "expected"],
               [(t, len(t.split())) for t in word_inputs],
               "assert word_count(text) == expected"),
            fn("test_truncate", ["text", "limit", "expected"],
               [(t, limit, ref_truncate(t, limit)) for t, limit in truncate_inputs],
               "assert truncate(text, limit) == expected"),
        ],
    )

    # ---- app/utils/timeutil.py : 40 ---------------------------------------
    block(
        "test_utils_timeutil.py",
        "app.utils.timeutil",
        "Duration parsing and formatting.",
        "from app.utils.timeutil import add_seconds, format_duration, is_business_day, parse_duration_ms",
        [
            fn("test_parse_duration_ms", ["text", "expected"],
               [(t, ref_parse_duration_ms(t)) for t in
                ["0ms", "1ms", "999ms", "1s", "2s", "59s", "1m", "90m", "1h", "24h"]],
               "assert parse_duration_ms(text) == expected"),
            fn("test_format_duration", ["milliseconds", "expected"],
               [(m, ref_format_duration(m)) for m in
                [0, 1, 500, 1000, 2000, 59000, 60000, 90000, 3600000, 7200000]],
               "assert format_duration(milliseconds) == expected"),
            fn("test_add_seconds", ["epoch_seconds", "seconds", "expected"],
               [(0, 0, 0), (0, 60, 60), (1000, -500, 500), (86399, 1, 86400),
                (1700000000, 3600, 1700003600), (1, 1, 2), (-10, 5, -5), (99, 1, 100),
                (123456789, 987654, 124444443), (5, 0, 5)],
               "assert add_seconds(epoch_seconds, seconds) == expected"),
            fn("test_is_business_day", ["weekday", "expected"],
               [(0, True), (1, True), (2, True), (3, True), (4, True), (5, False), (6, False),
                (-1, False), (7, False), (3, True)],
               "assert is_business_day(weekday) is expected"),
        ],
    )

    # ---- app/utils/ids.py : 40 --------------------------------------------
    numbers = [0, 1, 9, 10, 99, 100, 999, 1000, 999999, 123456]
    block(
        "test_utils_ids.py",
        "app.utils.ids",
        "Identifier formatting and sequencing.",
        "from app.utils.ids import account_id, invoice_id, is_valid_id, next_sequence",
        [
            fn("test_account_id", ["number", "expected"],
               [(n, f"ACCT-{n:06d}") for n in numbers],
               "assert account_id(number) == expected"),
            fn("test_invoice_id", ["number", "expected"],
               [(n, f"INV-{n:06d}") for n in numbers],
               "assert invoice_id(number) == expected"),
            fn("test_is_valid_id", ["value", "expected"],
               [("ACCT-000001", True), ("INV-000001", True), ("TXN-123456", True),
                ("ACCT-00000", False), ("ACCT-0000001", False), ("acct-000001", False),
                ("ACCT000001", False), ("ACCT-00000a", False), ("XXXX-000001", False), (None, False)],
               "assert is_valid_id(value) is expected"),
            fn("test_next_sequence", ["values", "expected"],
               [([], 1), (["ACCT-000001"], 2), (["ACCT-000009", "ACCT-000003"], 10),
                (["nonsense"], 1), (["INV-000100", "TXN-000050"], 101), (["ACCT-999999"], 1000000),
                (["ACCT-000001", "nonsense", "TXN-000002"], 3), (["ACCT-00000"], 1),
                (["ACCT-000020", "INV-000019"], 21), (["TXN-000042"], 43)],
               "assert next_sequence(values) == expected"),
        ],
    )

    # ---- app/api/auth.py : 20 ---------------------------------------------
    tokens = ["", "a", "token-1", "1" * 64, "üñïçø∂é"]
    block(
        "test_auth.py",
        "app.api.auth",
        "Token hashing and role scopes.",
        "from app.api.auth import hash_token, is_allowed, scopes_for, verify_token",
        [
            fn("test_hash_token", ["token", "expected"],
               [(t, ref_hash_token(t)) for t in tokens],
               "assert hash_token(token) == expected"),
            fn("test_verify_token", ["token", "digest", "expected"],
               [(tokens[0], ref_hash_token(tokens[0]), True),
                (tokens[1], ref_hash_token(tokens[1]), True),
                (tokens[2], ref_hash_token(tokens[2]), True),
                (tokens[3], ref_hash_token(tokens[3]), True),
                (tokens[4], ref_hash_token(tokens[4] + "x"), False)],
               "assert verify_token(token, digest) is expected"),
            fn("test_scopes_for", ["role", "expected"],
               [(r, ref_scopes_for(r)) for r in ["viewer", "operator", "admin", "auditor", "ADMIN"]],
               "assert scopes_for(role) == expected"),
            fn("test_is_allowed", ["role", "scope", "expected"],
               [("viewer", "read", True), ("viewer", "write", False), ("operator", "write", True),
                ("admin", "delete", True), ("auditor", "read", False)],
               "assert is_allowed(role, scope) is expected"),
        ],
    )

    # ---- app/utils/money.py : 32 ------------------------------------------
    block(
        "test_utils_money.py",
        "app.utils.money",
        "Integer-cent money arithmetic.",
        "from app.utils.money import apply_rate, format_amount, split_amount, to_cents",
        [
            fn("test_to_cents", ["amount_text", "expected"],
               [(t, ref_to_cents(t)) for t in
                ["0.00", "0.10", "1.00", "99.99", "-2.50", "+3", "7", "12.3"]],
               "assert to_cents(amount_text) == expected"),
            fn("test_format_amount", ["cents", "expected"],
               [(c, ref_format_amount(c)) for c in [0, 5, 10, 100, 9999, -250, 123456, 1]],
               "assert format_amount(cents) == expected"),
            fn("test_apply_rate", ["cents", "basis_points", "expected"],
               [(c, bps, ref_apply_rate(c, bps)) for c, bps in
                [(10000, 250), (10000, 0), (999, 1), (100, 10000), (12345, 725), (1, 9999),
                 (5000, 5000), (999999, 1)]],
               "assert apply_rate(cents, basis_points) == expected"),
            fn("test_split_amount", ["cents", "parts", "expected"],
               [(c, p, ref_split_amount(c, p)) for c, p in
                [(0, 1), (10, 1), (10, 2), (10, 3), (100, 3), (7, 4), (-10, 2), (1, 3)]],
               "assert split_amount(cents, parts) == expected"),
        ],
    )

    # ---- app/workers/report_worker.py : 6 (hand-written template) ----------
    block(
        "test_report_worker.py",
        "app.workers.report_worker",
        None,
        None,
        None,
        template="test_report_worker.py",
        rows=REPORT_WORKER_ROWS,
    )

    # ---- app/services/pricing_service.py : 56 ------------------------------
    skus = ["starter", "standard", "premium", "starter", "standard", "premium", "starter",
            "standard", "premium", "starter", "standard", "premium", "starter", "standard"]
    block(
        "test_pricing_service.py",
        "app.services.pricing_service",
        "Price tables, discounts and tiers.",
        "from app.services.pricing_service import base_price, discounted, tier_for, total_for",
        [
            fn("test_base_price", ["sku", "expected"],
               [(s, ref_base_price(s)) for s in skus],
               "assert base_price(sku) == expected"),
            fn("test_discounted", ["price_cents", "percent", "expected"],
               [(p, pct, ref_discounted(p, pct)) for p, pct in
                [(900, 0), (900, 10), (2500, 25), (7900, 50), (100, 100), (99, 33), (1, 50),
                 (123456, 7), (500, 1), (500, 99), (1000, 5), (1000, 15), (2500, 100), (7, 7)]],
               "assert discounted(price_cents, percent) == expected"),
            fn("test_tier_for", ["units", "expected"],
               [(u, ref_tier_for(u)) for u in
                [0, 1, 9, 10, 11, 99, 100, 101, 1000, 5, 50, 500, 99, 100]],
               "assert tier_for(units) == expected"),
            fn("test_total_for", ["sku", "units", "expected"],
               [(s, u, ref_base_price(s) * u) for s, u in
                [("starter", 0), ("starter", 1), ("starter", 10), ("standard", 1), ("standard", 3),
                 ("premium", 2), ("premium", 0), ("starter", 100), ("standard", 100), ("premium", 10),
                 ("starter", 7), ("standard", 11), ("premium", 13), ("starter", 99)]],
               "assert total_for(sku, units) == expected"),
        ],
    )

    # ---- app/services/discount_service.py : 30 -----------------------------
    block(
        "test_discount_service.py",
        "app.services.discount_service",
        "Discount code validation and application.",
        "from app.services.discount_service import apply, discount_for, is_valid_code",
        [
            fn("test_is_valid_code", ["code", "expected"],
               [(c, ref_discount_valid(c)) for c in
                ["SAVE10", "SAVE99", "SAVE00", "SAVE1", "SAVE100", "save10", "SAVE1A", "NOPE10", "", None]],
               "assert is_valid_code(code) is expected"),
            fn("test_discount_for", ["code", "cents", "expected"],
               [(c, cents, ref_discount_for(c, cents)) for c, cents in
                [("SAVE10", 1000), ("SAVE50", 999), ("SAVE00", 12345), ("SAVE99", 100),
                 ("NOPE10", 1000), ("SAVE20", 0), ("SAVE33", 1234), ("SAVE10", 1),
                 ("SAVE10", 9), ("SAVE99", 99999)]],
               "assert discount_for(code, cents) == expected"),
            fn("test_apply", ["code", "cents", "expected"],
               [(c, cents, (ref_discount_for(c, cents), cents - ref_discount_for(c, cents)))
                for c, cents in
                [("SAVE10", 1000), ("SAVE25", 2500), ("SAVE00", 0), ("SAVE50", 7),
                 ("NOPE10", 500), ("SAVE99", 10000), ("SAVE01", 99), ("SAVE10", 55),
                 ("SAVE20", 123456), ("SAVE75", 4)]],
               "assert apply(code, cents) == expected"),
        ],
    )

    # ---- app/workers/export_worker.py : 30 ---------------------------------
    block(
        "test_export_worker.py",
        "app.workers.export_worker",
        "Batch splitting for export jobs.",
        "from app.workers.export_worker import batch, batch_count, flatten",
        [
            fn("test_batch", ["items", "size", "expected"],
               [([], 3, []), ([1], 3, [[1]]), ([1, 2, 3], 3, [[1, 2, 3]]),
                ([1, 2, 3, 4], 3, [[1, 2, 3], [4]]),
                (list(range(10)), 4, [list(range(0, 4)), list(range(4, 8)), [8, 9]]),
                ([1, 2], 1, [[1], [2]]), ([], 1, []),
                (list(range(6)), 2, [[0, 1], [2, 3], [4, 5]]),
                ([1, 2, 3], 10, [[1, 2, 3]]), (list(range(5)), 5, [[0, 1, 2, 3, 4]])],
               "assert batch(items, size) == expected"),
            fn("test_batch_count", ["total", "size", "expected"],
               [(0, 10, 0), (1, 10, 1), (10, 10, 1), (11, 10, 2), (100, 100, 1),
                (101, 100, 2), (7, 3, 3), (9, 3, 3), (10, 3, 4), (1000, 7, 143)],
               "assert batch_count(total, size) == expected"),
            fn("test_flatten", ["batches", "expected"],
               [([], []), ([[1]], [1]), ([[1, 2], [3]], [1, 2, 3]), ([[], [1]], [1]),
                ([[1], [], [2]], [1, 2]), ([[0, 1, 2], [3, 4]], [0, 1, 2, 3, 4]),
                ([["a"], ["b", "c"]], ["a", "b", "c"]), ([[1, 2, 3]], [1, 2, 3]),
                ([[1], [2], [3], [4]], [1, 2, 3, 4]), ([[], []], [])],
               "assert flatten(batches) == expected"),
        ],
    )

    # ---- app/services/cache_service.py : 2 (shape only, never format) ------
    block(
        "test_cache_service.py",
        "app.services.cache_service",
        "Cache writing: asserts shape only, never the wire format.\nThe format itself is pinned end-to-end by tests/test_report_worker.py.",
        "from app.services.cache_service import read_cache_entries, write_cache_entry",
        [
            fn("test_write_cache_entry_returns_a_payload", ["key", "status", "count"],
               [("key-0", "ok", 0)],
               "sink = []\n"
               "payload = write_cache_entry(key, status, count, sink=sink)\n"
               "assert isinstance(payload, str) and payload\n"
               "assert sink == [payload]"),
            fn("test_read_cache_entries", ["blank"],
               [(None,)],
               "path = tmp_path / 'cache.txt'\n"
               "path.write_text('a\\n\\n b \\n', encoding='utf-8')\n"
               "assert read_cache_entries(path) == ['a', ' b ']",
               fixtures=["tmp_path"]),
        ],
    )

    # ---- app/services/account_service.py : 2 -------------------------------
    block(
        "test_account_service.py",
        "app.services.account_service",
        "Account lifecycle, asserting shape only.",
        "from app.services.account_service import account_summary, open_account, record_event",
        [
            fn("test_open_account", ["account_id", "owner"],
               [("ACCT-000001", "owner-1")],
               "assert open_account(account_id, owner) == {'id': account_id, 'owner': owner, 'state': 'open'}"),
            fn("test_account_summary_counts_cached_entries", ["blank"],
               [(None,)],
               "path = tmp_path / 'cache.txt'\n"
               "path.write_text('ACCT-000001|x|1\\nACCT-000002|x|1\\n', encoding='utf-8')\n"
               "assert record_event('ACCT-000001', 'ok') is not None\n"
               "assert account_summary('ACCT-000001', path) == 1",
               fixtures=["tmp_path"]),
        ],
    )

    # ---- app/services/search_service.py : 2 --------------------------------
    block(
        "test_search_service.py",
        "app.services.search_service",
        "Search and ranking over entry payloads.",
        "from app.services.search_service import match_accounts, rank",
        [
            fn("test_match_accounts", ["entries", "query"],
               [(["Alpha", "beta", "ALPHABET", "gamma"], "alph")],
               "assert match_accounts(entries, query) == ['Alpha', 'ALPHABET']"),
            fn("test_rank", ["entries", "query"],
               [(["bb", "a", "bbb"], "b")],
               "assert rank(entries, query) == ['bbb', 'bb']"),
        ],
    )

    # ---- app/services/export_service.py : 2 --------------------------------
    block(
        "test_export_service.py",
        "app.services.export_service",
        "CSV rendering of cached entries.",
        "from app.services.export_service import export_filename, export_rows",
        [
            fn("test_export_rows", ["entries"],
               [(["ab", "c"],)],
               "assert export_rows(entries) == ['entry,length', 'ab,2', 'c,1']"),
            fn("test_export_filename", ["prefix", "stamp"],
               [("recent", "20260101")],
               "assert export_filename(prefix, stamp) == 'recent-20260101.csv'"),
        ],
    )

    # ---- app/api/routes.py : 2 ---------------------------------------------
    block(
        "test_routes.py",
        "app.api.routes",
        "Route resolution.",
        "from app.api.routes import is_known, resolve",
        [
            fn("test_resolve", ["method", "path"],
               [("GET", "/accounts")],
               "assert resolve(method, path) == 'list_accounts'"),
            fn("test_is_known", ["method", "path"],
               [("GET", "/reports/recent")],
               "assert is_known(method, path) is True"),
        ],
    )

    # ---- app/api/serializers.py : 2 ----------------------------------------
    block(
        "test_serializers.py",
        "app.api.serializers",
        "Response serialisation.",
        "from app.api.serializers import serialise_account, serialise_error",
        [
            fn("test_serialise_account", ["account"],
               [({"id": "ACCT-1", "state": "open"},)],
               "assert serialise_account(account) == 'id=ACCT-1;state=open'"),
            fn("test_serialise_error", ["code", "message"],
               [(404, "missing")],
               "assert serialise_error(code, message) == 'error 404: missing'"),
        ],
    )

    # ---- app/workers/ingest_worker.py : 3 ----------------------------------
    block(
        "test_ingest_worker.py",
        "app.workers.ingest_worker",
        "Event ingestion. Payloads are asserted by shape only, never by wire\nformat, so these tests stay green across format changes.",
        "from app.workers.ingest_worker import ingest_batch, normalise_status, summarise_batch",
        [
            fn("test_normalise_status", ["raw"],
               [("",)],
               "assert normalise_status(raw) == 'retry'"),
            fn("test_ingest_batch", ["events"],
               [([{"key": "k1", "status": "ok"}, {"key": "k2", "status": "nope"}],)],
               "payloads = ingest_batch(events)\n"
               "assert len(payloads) == 2\n"
               "assert all(isinstance(payload, str) and payload for payload in payloads)"),
            fn("test_summarise_batch", ["events"],
               [([{"key": "a", "status": "ok"}, {"key": "b", "status": "weird"}],)],
               "assert summarise_batch(events) == {'ok': 1, 'retry': 1}"),
        ],
    )

    # ---- app/workers/cleanup_worker.py : 2 ---------------------------------
    block(
        "test_cleanup_worker.py",
        "app.workers.cleanup_worker",
        "Cache trimming.",
        "from app.workers.cleanup_worker import clean_cache, trim",
        [
            fn("test_trim", ["entries", "limit", "expected"],
               [([1, 2, 3], 2, [2, 3])],
               "assert trim(entries, limit) == expected"),
            fn("test_clean_cache", ["limit", "expected"],
               [(1, (1, 1))],
               "path = tmp_path / 'cache.txt'\n"
               "path.write_text('one\\n\\n two \\n', encoding='utf-8')\n"
               "assert clean_cache(path, limit) == expected",
               fixtures=["tmp_path"]),
        ],
    )

    # ---- app/services/audit_service.py : 1 ---------------------------------
    block(
        "test_audit_service.py",
        "app.services.audit_service",
        "Audit sensitivity list.",
        "from app.services.audit_service import audit_line, is_sensitive",
        [
            fn("test_is_sensitive", ["action", "expected"],
               [("delete", True)],
               "assert is_sensitive(action) is expected\n"
               "assert audit_line(action, 'actor') == f'actor {action}'"),
        ],
    )

    # ---- app/services/notification_service.py : 20 -------------------------
    block(
        "test_notification_service.py",
        "app.services.notification_service",
        "Notification channel selection and rendering.",
        "from app.services.notification_service import body_length, channels_for, render_message, should_notify",
        [
            fn("test_channels_for", ["account", "expected"],
               [({"state": "open"}, ["email", "push"]), ({"state": "closed"}, ["email"]),
                ({"state": "frozen"}, ["email"]), ({}, ["email"]),
                ({"state": "open", "id": "ACCT-1"}, ["email", "push"])],
               "assert channels_for(account) == expected"),
            fn("test_render_message", ["template", "values", "expected"],
               [("hello {name}", {"name": "world"}, "hello world"),
                ("{a}-{b}", {"a": "1", "b": "2"}, "1-2"),
                ("no placeholders", {}, "no placeholders"),
                ("{x}{x}", {"x": "ab"}, "abab"),
                ("total {amount} cents", {"amount": 250}, "total 250 cents")],
               "assert render_message(template, values) == expected"),
            fn("test_should_notify", ["account", "event", "expected"],
               [({"state": "open"}, "settled", True), ({"state": "closed"}, "settled", True),
                ({"state": "open"}, "opened", True), ({"state": "closed"}, "opened", False),
                ({"state": "open"}, "ignored", False)],
               "assert should_notify(account, event) is expected"),
            fn("test_body_length", ["message", "expected"],
               [("", 0), ("abc", 3), ("x" * 239, 239), ("x" * 240, 240), ("x" * 241, 240)],
               "assert body_length(message) == expected"),
        ],
    )

    # ---- app/services/invoice_service.py : 3 -------------------------------
    block(
        "test_invoice_service.py",
        "app.services.invoice_service",
        "Invoice issuing and due dates.",
        "from app.services.invoice_service import due_date, is_overdue, issue",
        [
            fn("test_issue", ["account", "invoice"],
               [({"id": "ACCT-1"}, {"id": "INV-2"})],
               "assert issue(account, invoice) == 'invoice INV-2 issued to ACCT-1'"),
            fn("test_due_date", ["issued_on_day"],
               [(0,)],
               "assert due_date(issued_on_day) == 30"),
            fn("test_is_overdue", ["due_day", "today"],
               [(30, 31)],
               "assert is_overdue(due_day, today) is True"),
        ],
    )

    return blocks


# --------------------------------------------------------------------------- #
# rendering                                                                    #
# --------------------------------------------------------------------------- #
def render_block(block_spec, start_id):
    """Render one test file; return (source, [(test_id, name, description, runtime)])."""
    lines = ['"""' + block_spec["doc"] + '"""', "", "import pytest", ""]
    lines.append(block_spec["imports"])
    lines.append("")
    lines.append("")
    rows = []
    next_id = start_id
    for function in block_spec["functions"]:
        params_src = ", ".join(function["params"])
        lines.append("@pytest.mark.parametrize(")
        lines.append(f"    {params_src!r},")
        lines.append("    [")
        for case in function["cases"]:
            test_id = f"T-{next_id:04d}"
            next_id += 1
            rendered = ", ".join(repr(value) for value in case)
            lines.append(f'        pytest.param({rendered}, id="{test_id}"),')
            description = function["name"]
            if case:
                description = f"{function['name']} with {', '.join(repr(v) for v in case[:2])}"
            rows.append((test_id, function["name"], description[:90], 5 + (int(test_id[2:]) * 7) % 60))
        lines.append("    ],")
        lines.append(")")
        signature = function["params"] + function.get("fixtures", [])
        lines.append(f"def {function['name']}({', '.join(signature)}):")
        lines.append(f'    """{function["name"].replace("test_", "").replace("_", " ").capitalize()}."""')
        for body_line in function["body"].splitlines():
            lines.append(f"    {body_line}" if body_line else "")
        lines.append("")
        lines.append("")
    return "\n".join(lines) + "\n", rows


def main():
    blocks = generate_all()
    TESTS.mkdir(parents=True, exist_ok=True)
    inventory_rows = []
    next_id = 1
    for block_spec in blocks:
        if block_spec["template"]:
            source = (pathlib.Path(__file__).resolve().parent / "templates" / block_spec["template"]).read_text(
                encoding="utf-8"
            )
            (TESTS / block_spec["file"]).write_text(source, encoding="utf-8")
            rows = list(block_spec["rows"])
        else:
            source, rows = render_block(block_spec, next_id)
            (TESTS / block_spec["file"]).write_text(source, encoding="utf-8")
        for test_id, name, description, runtime in rows:
            inventory_rows.append(
                {
                    "test_id": test_id,
                    "test_name": name,
                    "module": block_spec["module"],
                    "description": description,
                    "avg_runtime_ms": runtime,
                }
            )
        next_id += len(rows)
    total = len(inventory_rows)
    if total != EXPECTED_TOTAL:
        raise SystemExit(f"fixture drift: generated {total} tests, expected {EXPECTED_TOTAL}")
    write_inventory(inventory_rows)
    print(f"generated {total} tests across {len(blocks)} files; inventory rewritten")


def write_inventory(rows):
    """Write inventory.csv and inventory.xlsx from the same rows."""
    header = ["test_id", "test_name", "module", "description", "avg_runtime_ms", "status", "owner"]
    # The link-correction row: the inventory CLAIMS app.utils.ids for the first
    # text test while the test file actually imports app.utils.text. The declared
    # link is kept as evidence; the derived link is used (defined behaviour).
    for row in rows:
        if row["test_id"] == "T-0146":
            row["module"] = "app.utils.ids"
    with (ROOT / "inventory.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        for row in rows:
            writer.writerow([row["test_id"], row["test_name"], row["module"], row["description"],
                             row["avg_runtime_ms"], "active", "qa-team"])
    try:
        from openpyxl import Workbook
    except ImportError:
        print("openpyxl unavailable: inventory.xlsx not written")
        return
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "inventory"
    sheet.append(header)
    for row in rows:
        sheet.append([row["test_id"], row["test_name"], row["module"], row["description"],
                      row["avg_runtime_ms"], "active", "qa-team"])
    workbook.save(ROOT / "inventory.xlsx")


if __name__ == "__main__":
    main()
